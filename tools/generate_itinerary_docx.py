from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('output/family_trip_berlin_eberswalde_luebeck_hamburg_links.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_hyperlink(paragraph, text, url, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    rPr.append(underline)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    new_run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_bullet(doc, pieces, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    set_rtl(p)
    for kind, text, value in pieces:
        if kind == 'text':
            r = p.add_run(text)
            r.font.name = 'Arial'
        elif kind == 'bold':
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Arial'
        elif kind == 'link':
            add_hyperlink(p, text, value)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_rtl(p)
    return p


def add_para(doc, text='', bold_prefix=None):
    p = doc.add_paragraph()
    set_rtl(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(10.5)
for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    styles[name].font.name = 'Arial'
styles['Title'].font.color.rgb = RGBColor(31, 78, 121)
styles['Heading 1'].font.color.rgb = RGBColor(31, 78, 121)
styles['Heading 2'].font.color.rgb = RGBColor(54, 95, 145)

p = doc.add_paragraph(style='Title')
set_rtl(p)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('חופשה משפחתית של רני – מסלול מעודכן עם קישורים')
p2 = doc.add_paragraph()
set_rtl(p2)
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run('12–24 באוגוסט 2026 | ברלין – אברסוולדה | סבא וסבתא ממשיכים לליבק ולהמבורג')
r.bold = True

add_heading(doc, 'תקציר המסלול', 1)
add_para(doc, 'ברלין 12–16.8 בשתי דירות Airbnb; ב־16.8 איסוף שתי מכוניות ונסיעה לאברסוולדה; לינה באברסוולדה עד 20.8. ב־20.8 הסבא והסבתא ממשיכים ברכב אחד לליבק, ויתר המשפחה חוזרים ברכב השני לברלין.')
add_para(doc, 'הרכב הקבוצה: תינוק בן שנה וחצי, ארבעה צעירים בני כ־30, וסבא וסבתא בני כ־60.')

add_heading(doc, 'טיסות', 1)
add_bullet(doc, [('bold','12.8.2026: ',None),('text','תל אביב → ברלין, LY2371, נחיתה ב־BER בשעה 09:35. ',None),('link','אתר אל על','https://www.elal.com/'),('text',' | ',None),('link','מפת BER','https://www.google.com/maps/search/?api=1&query=Berlin+Brandenburg+Airport')])
add_bullet(doc, [('bold','24.8.2026: ',None),('text','המבורג → אתונה → תל אביב, A3825 + A3926, המראה 11:05. ',None),('link','אתר Aegean','https://en.aegeanair.com/'),('text',' | ',None),('link','מפת HAM','https://www.google.com/maps/search/?api=1&query=Hamburg+Airport')])

add_heading(doc, 'לינות, כתובות וקישורי הזמנה', 1)
add_heading(doc, 'ברלין | 12–16.8', 2)
add_bullet(doc, [('bold','דירה מרכזית: ',None),('text','Gleimstraße 53, 10437 Berlin — ',None),('link','ניווט ב־Google Maps','https://www.google.com/maps/search/?api=1&query=Gleimstra%C3%9Fe+53%2C+10437+Berlin%2C+Germany')])
add_bullet(doc, [('bold','דירה משנית: ',None),('text','Kopenhagener Straße 45, 10437 Berlin — ',None),('link','ניווט ב־Google Maps','https://www.google.com/maps/search/?api=1&query=Kopenhagener+Stra%C3%9Fe+45%2C+10437+Berlin%2C+Germany')])
add_bullet(doc, [('link','פתיחת הזמנת Airbnb בברלין','https://www.airbnb.com/l/i0ngSu0p?s=67&unique_share_id=9d7f6b50-79fa-4a82-939d-11c9382235d9')])
add_heading(doc, 'אברסוולדה | 16–20.8', 2)
add_bullet(doc, [('bold','כתובת הבית: ',None),('text','Eichwerder Straße, 16225 Eberswalde, Brandenburg — ',None),('link','ניווט ב־Google Maps','https://www.google.com/maps/search/?api=1&query=Eichwerder+Stra%C3%9Fe%2C+16225+Eberswalde%2C+Germany')])
add_bullet(doc, [('link','פתיחת הזמנת Airbnb באברסוולדה','https://www.airbnb.com/l/pDhJRKwE?s=67&unique_share_id=85454bbe-ec22-45fb-bfaf-44754bd42be2')])
add_heading(doc, 'ליבק | 20–21.8 — סבא וסבתא בלבד', 2)
add_bullet(doc, [('bold','Radisson Blu Senator Hotel: ',None),('text','Willy-Brandt-Allee 6, 23554 Lübeck — ',None),('link','ניווט','https://www.google.com/maps/search/?api=1&query=Radisson+Blu+Senator+Hotel+L%C3%BCbeck'),('text',' | ',None),('link','אתר המלון','https://www.radissonhotels.com/en-us/hotels/radisson-blu-luebeck-senator')])
add_bullet(doc, [('text','הזמנה 2YG382ZY, חדר Standard Room עם מיטת King וארוחת בוקר.',None)])
add_heading(doc, 'המבורג | 21–24.8 — סבא וסבתא בלבד', 2)
add_bullet(doc, [('bold','Reichshof Hotel Hamburg: ',None),('text','Kirchenallee 34–36, 20099 Hamburg — ',None),('link','ניווט','https://www.google.com/maps/search/?api=1&query=Reichshof+Hotel+Hamburg'),('text',' | ',None),('link','אתר המלון','https://www.reichshof-hotel-hamburg.de/')])

add_heading(doc, 'רכבים', 1)
add_bullet(doc, [('text','שתי מכוניות נאספות בברלין ב־16.8. יש לוודא כיסא תינוק, נהגים מורשים ומקום למזוודות. ',None),('link','Europcar','https://www.europcar.com/'),('text',' | ',None),('link','EconomyCarRentals','https://www.economycarrentals.com/')])
add_bullet(doc, [('text','ב־20.8: רכב אחד עם הסבא והסבתא ממשיך לליבק ומוחזר שם ב־21.8; הרכב השני חוזר עם יתר המשפחה לברלין ומוחזר בברלין ב־20.8.',None)])

doc.add_page_break()
add_heading(doc, 'ברלין | 12–16.8', 1)
add_heading(doc, 'יום רביעי 12.8 — הגעה והתאקלמות', 2)
add_bullet(doc, [('text','09:35 נחיתה ב־BER, נסיעה לדירות, השארת מזוודות ומנוחה.',None)])
add_bullet(doc, [('text','אחר הצהריים: טיול קל ב־Prenzlauer Berg, Hackescher Markt או סביב אי המוזיאונים.',None)])
add_bullet(doc, [('bold','18:00 — Song Nguu, הזמנה ל־6 סועדים. ',None),('text','Am Falkpl. 5, 10437 Berlin — ',None),('link','ניווט למסעדה','https://www.google.com/maps/search/?api=1&query=Song+Nguu%2C+Am+Falkpl.+5%2C+10437+Berlin')])

add_heading(doc, 'יום חמישי 13.8 — אי המוזיאונים וסיור אדריכלות', 2)
add_bullet(doc, [('link','אי המוזיאונים','https://www.museumsinsel-berlin.de/en/'),('text',', ',None),('link','קתדרלת ברלין','https://www.berlinerdom.de/en/'),('text',' ו־Lustgarten.',None)])
add_bullet(doc, [('bold','14:30 — Architecture Tour עם Natalie. ',None),('text','מפגש בכניסה ל־',None),('link','James-Simon-Galerie','https://www.smb.museum/en/museums-institutions/james-simon-galerie/home/'),('text',' — ',None),('link','מפה','https://www.google.com/maps/search/?api=1&query=James-Simon-Galerie+Berlin'),('text',' | ',None),('link','Tours of Berlin','https://toursofberlin.com/private-tours/')])

add_heading(doc, 'יום שישי 14.8 — Off the Beaten Path', 2)
add_bullet(doc, [('text','בוקר: Nikolaiviertel, הליכה לאורך ה־Spree, או ',None),('link','East Side Gallery','https://www.eastsidegalleryexhibition.com/')])
add_bullet(doc, [('bold','14:00 — סיור עם Martin. ',None),('text','מפגש בתחנת U-Bahn Klosterstraße (U2) — ',None),('link','מפת נקודת המפגש','https://www.google.com/maps/search/?api=1&query=Klosterstra%C3%9Fe+U-Bahn+Berlin'),('text',' | ',None),('link','פרטי הסיור','https://toursofberlin.com/private-tours/')])
add_bullet(doc, [('text','חלופה מקורה: ',None),('link','DDR Museum','https://www.ddr-museum.de/en')])

add_heading(doc, 'יום שבת 15.8 — יום נוסף בברלין', 2)
add_bullet(doc, [('link','הזמנת כניסה לכיפת הרייכסטאג','https://www.bundestag.de/en/visittheBundestag/dome/registration-245686'),('text',', שער ברנדנבורג ו־Tiergarten.',None)])
add_bullet(doc, [('link','שיט קצר על ה־Spree','https://www.visitberlin.de/en/boat-tours'),('text',' או ביקור ב־DDR Museum.',None)])
add_bullet(doc, [('text','חלופה לצעירים: ',None),('link','Tempelhofer Feld','https://www.thf-berlin.de/en/'),('text',', Kreuzberg או Friedrichshain.',None)])
add_bullet(doc, [('bold','חשוב: ',None),('text','לקנות בערב מצרכים ליום ראשון ולבית באברסוולדה, משום שרוב הסופרמרקטים סגורים ביום ראשון.',None)])

add_heading(doc, 'יום ראשון 16.8 — רכבים, נסיעה ואפשרות לקונצרט', 2)
add_bullet(doc, [('text','איסוף שתי המכוניות בברלין ונסיעה לאברסוולדה.',None)])
add_bullet(doc, [('bold','15:00 — Choriner Musiksommer במנזר Chorin. ',None),('link','אתר הפסטיבל וכרטיסים','https://www.choriner-musiksommer.de/'),('text',' | ',None),('link','אתר מנזר Chorin','https://www.kloster-chorin.org/'),('text',' | ',None),('link','ניווט','https://www.google.com/maps/search/?api=1&query=Kloster+Chorin')])
add_bullet(doc, [('text','בגלל המעבר והתינוק אפשר להתפצל: רכב אחד לקונצרט ורכב אחד ישירות ל־Airbnb.',None)])

doc.add_page_break()
add_heading(doc, 'אברסוולדה | 16–20.8', 1)
add_para(doc, 'המסלול באזור בנוי סביב אתר מרכזי אחד ביום, עם אפשרות להתפצל בין שתי המכוניות.')
add_heading(doc, 'יום שני 17.8 — Werbellinsee', 2)
add_bullet(doc, [('link','Werbellinsee — מפה וניווט','https://www.google.com/maps/search/?api=1&query=Werbellinsee+Germany'),('text','. יום אגם: רחצה, פיקניק, מנוחה בצל והליכה קלה.',None)])
add_heading(doc, 'יום שלישי 18.8 — Chorin, Brodowin ו־Parsteiner See', 2)
add_bullet(doc, [('link','Kloster Chorin','https://www.kloster-chorin.org/'),('text',' | ',None),('link','Brodowin','https://www.google.com/maps/search/?api=1&query=Brodowin+Germany'),('text',' | ',None),('link','Parsteiner See','https://www.google.com/maps/search/?api=1&query=Parsteiner+See')])
add_bullet(doc, [('text','חלופה קרובה: ',None),('link','Eberswalde Zoo','https://www.zoo.eberswalde.de/')])
add_heading(doc, 'יום רביעי 19.8 — יום גמיש', 2)
add_bullet(doc, [('link','Wandlitzsee','https://www.google.com/maps/search/?api=1&query=Wandlitzsee'),('text',' — חוף מסודר ותשתיות למשפחות.',None)])
add_bullet(doc, [('link','Niederfinow Schiffshebewerk','https://www.schiffshebewerk-niederfinow.com/'),('text',' — ',None),('link','ניווט','https://www.google.com/maps/search/?api=1&query=Schiffshebewerk+Niederfinow')])
add_bullet(doc, [('text','חלופות קלות או ליום גשם: ',None),('link','Familiengarten Eberswalde','https://www.familiengarten-eberswalde.de/'),('text',', ',None),('link','Eberswalde Zoo','https://www.zoo.eberswalde.de/'),('text',' או ',None),('link','Freizeitbad baff','https://www.baff-bad.de/')])
add_heading(doc, 'יום חמישי 20.8 — פיצול המשפחה', 2)
add_bullet(doc, [('text','סבא וסבתא: נסיעה לליבק, צ׳ק־אין במלון וערב בעיר העתיקה.',None)])
add_bullet(doc, [('text','יתר המשפחה: נסיעה לברלין והחזרת הרכב באותו יום.',None)])
add_bullet(doc, [('bold','19:30 — Nordic Sounds בליבק. ',None),('link','אתר Schleswig-Holstein Musik Festival וכרטיסים','https://www.shmf.de/en')])

add_heading(doc, 'ליבק | 20–21.8 — סבא וסבתא', 1)
add_bullet(doc, [('link','אתר התיירות הרשמי של ליבק','https://www.visit-luebeck.com/')])
add_bullet(doc, [('link','Holstentor','https://museum-holstentor.de/'),('text',' | ',None),('link','Marienkirche','https://www.st-marien-luebeck.de/'),('text',' | ',None),('link','Niederegger','https://www.niederegger.de/'),('text',' | ',None),('link','European Hansemuseum','https://www.hansemuseum.eu/'),('text',' | ',None),('link','Travemünde','https://www.visit-travemuende.com/')])

doc.add_page_break()
add_heading(doc, 'המבורג | 21–24.8 — סבא וסבתא', 1)
add_heading(doc, 'יום שישי 21.8 — הגעה וקונצרט', 2)
add_bullet(doc, [('text','החזרת הרכב בליבק והמשך להמבורג. טיול ב־',None),('link','Speicherstadt','https://www.hamburg.com/visitors/sights/architecture/11748624/speicherstadt/'),('text',' וב־',None),('link','HafenCity','https://www.hafencity.com/')])
add_bullet(doc, [('bold','20:00 — John Scofield’s Long Days Quartet באלבפילהרמוני. ',None),('link','עמוד הקונצרט והזמנת כרטיסים','https://www.elbphilharmonie.de/en/whats-on/john-scofield/27966'),('text',' | ',None),('link','ניווט','https://www.google.com/maps/search/?api=1&query=Elbphilharmonie+Hamburg'),('text',' | ',None),('link','מידע על ה־Plaza','https://www.elbphilharmonie.de/en/plaza')])
add_bullet(doc, [('text','להגיע לפחות 30 דקות מראש.',None)])
add_heading(doc, 'יום שבת 22.8 — מרכז היסטורי', 2)
add_bullet(doc, [('bold','11:00–13:00 — Historic Centre Free Walking Tour. ',None),('link','Robin and the Tourguides','https://www.robinandthetourguides.de/'),('text',' | ',None),('link','מפת Rathaus','https://www.google.com/maps/search/?api=1&query=Hamburg+Rathaus')])
add_bullet(doc, [('text','אחר הצהריים: Alster, ',None),('link','Planten un Blomen','https://www.hamburg.com/visitors/explore/green-hamburg/11875242/planten-un-blomen/'),('text',', ',None),('link','St. Michaelis','https://www.st-michaelis.de/en/'),('text',', Jungfernstieg ו־Mönckebergstraße.',None)])
add_bullet(doc, [('text','אפשרות נוספת: ',None),('link','Miniatur Wunderland — מומלץ להזמין מראש','https://www.miniatur-wunderland.com/')])
add_heading(doc, 'יום ראשון 23.8 — נמל וסנט פאולי', 2)
add_bullet(doc, [('link','נמל המבורג','https://www.hamburg.com/visitors/explore/port-of-hamburg/'),('text',', Fischmarkt או שיט מ־Landungsbrücken.',None)])
add_bullet(doc, [('bold','14:00–16:00 — Harbour & St. Pauli Tour. ',None),('link','פרטי הסיור','https://www.robinandthetourguides.de/')])
add_heading(doc, 'יום שני 24.8 — טיסה', 2)
add_bullet(doc, [('text','עזיבת המלון בסביבות 09:00 ונסיעה לנמל התעופה HAM.',None)])

add_heading(doc, 'הזמנות שכדאי לבצע מראש', 1)
add_bullet(doc, [('link','John Scofield, 21.8 בשעה 20:00','https://www.elbphilharmonie.de/en/whats-on/john-scofield/27966')])
add_bullet(doc, [('link','Nordic Sounds, ליבק, 20.8 בשעה 19:30','https://www.shmf.de/en')])
add_bullet(doc, [('link','Choriner Musiksommer, 16.8 בשעה 15:00','https://www.choriner-musiksommer.de/')])
add_bullet(doc, [('link','כיפת הרייכסטאג','https://www.bundestag.de/en/visittheBundestag/dome/registration-245686')])
add_bullet(doc, [('link','Miniatur Wunderland','https://www.miniatur-wunderland.com/'),('text',' ו/או ',None),('link','שיט בברלין','https://www.visitberlin.de/en/boat-tours')])

add_heading(doc, 'הערות פרקטיות', 1)
for text in [
    'לא לקבוע יותר מאטרקציה מחייבת אחת ביום.',
    'בימי אגמים להגיע מוקדם, למצוא צל ולשמור זמן לשנת הצהריים של התינוק.',
    'להחזיק ברכב תיק קבוע עם חיתולים, בגדי החלפה, מים, נשנושים, מגבת, עזרה ראשונה וצעצוע קטן.',
    'לפני כל יציאה לאגם לבדוק מזג אוויר, איכות מים, חניה, שירותים וצל.',
    'לעשות ביטוח נסיעות ולוודא כיסוי לרכב שכור ולנהגים הנוספים.'
]:
    add_bullet(doc, [('text',text,None)])

# Add footer with document note
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    set_rtl(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run('מסלול משפחתי מעודכן – יולי 2026 | הקישורים במסמך לחיצים')
    rr.font.size = Pt(8)
    rr.font.color.rgb = RGBColor(100,100,100)

# Core metadata
doc.core_properties.title = 'חופשה משפחתית של רני – מסלול מעודכן עם קישורים'
doc.core_properties.subject = 'ברלין, אברסוולדה, ליבק והמבורג – אוגוסט 2026'
doc.core_properties.author = 'Noam Kolodny'

doc.save(OUT)
print(OUT)
