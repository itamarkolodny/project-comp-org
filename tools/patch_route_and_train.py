from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX = Path('output/family_trip_berlin_eberswalde_luebeck_hamburg_links.docx')


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def clear_paragraph(paragraph):
    pPr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not pPr:
            paragraph._p.remove(child)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    rPr.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def fill_paragraph(paragraph, pieces):
    clear_paragraph(paragraph)
    set_rtl(paragraph)
    for kind, text, value in pieces:
        if kind == 'text':
            paragraph.add_run(text)
        elif kind == 'bold':
            r = paragraph.add_run(text)
            r.bold = True
        elif kind == 'link':
            add_hyperlink(paragraph, text, value)


def insert_after(paragraph, pieces, style='List Bullet'):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    fill_paragraph(new_para, pieces)
    return new_para


doc = Document(DOCX)

# Update the overview under rental cars.
for p in doc.paragraphs:
    if 'ב־20.8: רכב אחד עם הסבא והסבתא ממשיך לליבק' in p.text:
        fill_paragraph(p, [
            ('text', 'ב־20.8: הסבא והסבתא נוסעים ברכב מאברסוולדה לליבק דרך שוורין, עם עצירת צהריים וטיול קצר בעיר. הרכב מוחזר בליבק בבוקר 21.8. יתר המשפחה חוזרים ברכב השני לברלין ומחזירים אותו ב־20.8.', None)
        ])
        break

# Replace and expand the grandparents' driving day.
for p in doc.paragraphs:
    if p.text.startswith('סבא וסבתא: נסיעה לליבק'):
        fill_paragraph(p, [
            ('bold', 'סבא וסבתא — נסיעה לליבק דרך שוורין: ', None),
            ('text', 'יציאה מאברסוולדה בסביבות 09:00 ונסיעה מערבה דרך A10/A24/A14. עצירת צהריים בשוורין ולאחריה המשך לליבק.', None),
        ])
        anchor = p
        anchor = insert_after(anchor, [
            ('bold', 'מסלול נהיגה מוצע: ', None),
            ('link', 'Eberswalde → Schwerin → Lübeck ב־Google Maps', 'https://www.google.com/maps/dir/?api=1&origin=Eichwerder+Stra%C3%9Fe%2C+16225+Eberswalde%2C+Germany&destination=Radisson+Blu+Senator+Hotel+L%C3%BCbeck%2C+Willy-Brandt-Allee+6%2C+23554+L%C3%BCbeck%2C+Germany&waypoints=Schwerin+Castle%2C+Lenn%C3%A9stra%C3%9Fe+1%2C+19053+Schwerin%2C+Germany&travelmode=driving'),
            ('text', '. יש לבדוק עומסי תנועה בזמן אמת לפני היציאה.', None),
        ])
        anchor = insert_after(anchor, [
            ('bold', 'עצירת צהריים בשוורין (כשעתיים): ', None),
            ('text', 'חניה באזור העיר העתיקה, הליכה קצרה אל ', None),
            ('link', 'טירת שוורין', 'https://www.schwerin.de/en/visit-schwerin/attractions/schwerin-castle/'),
            ('text', ' ואל Markt. לארוחה מומלץ ', None),
            ('link', 'Weinhaus Uhle', 'https://www.weinhaus-uhle.de/'),
            ('text', ' באווירה היסטורית; לחלופה קלה יותר: ', None),
            ('link', 'Café Prag', 'https://www.restaurant-cafe-prag.de/'),
            ('text', '. מומלץ להזמין שולחן מראש.', None),
        ])
        anchor = insert_after(anchor, [
            ('bold', 'לוח זמנים משוער: ', None),
            ('text', 'יציאה 09:00; הגעה לשוורין סביב הצהריים; ארוחה והליכה קצרה 12:00–14:15; יציאה לליבק והגעה משוערת 15:30–16:00, בהתאם לתנועה.', None),
        ])
        break

# Replace the Lübeck-to-Hamburg driving wording with the train plan.
for p in doc.paragraphs:
    if p.text.startswith('החזרת הרכב בליבק והמשך להמבורג'):
        fill_paragraph(p, [
            ('bold', 'בוקר 21.8 — מעבר מליבק להמבורג ברכבת: ', None),
            ('text', 'החזרת הרכב בליבק, מעבר אל Lübeck Hbf ונסיעה ברכבת אזורית ישירה RE8 או RE80 אל Hamburg Hbf.', None),
        ])
        anchor = p
        anchor = insert_after(anchor, [
            ('text', 'הרכבות הישירות פועלות בדרך כלל בתדירות גבוהה ביום חול, וזמן הנסיעה הוא בערך 45–50 דקות. יש לבדוק את השעה המדויקת ואת עבודות המסילה סמוך לנסיעה ב־', None),
            ('link', 'Deutsche Bahn', 'https://int.bahn.de/en'),
            ('text', '.', None),
        ])
        anchor = insert_after(anchor, [
            ('text', 'אין צורך ברכב בהמבורג: מלון Reichshof נמצא מול Hamburg Hbf. ', None),
            ('link', 'מסלול הליכה מהתחנה למלון', 'https://www.google.com/maps/dir/?api=1&origin=Hamburg+Hauptbahnhof&destination=Reichshof+Hotel+Hamburg%2C+Kirchenallee+34-36%2C+20099+Hamburg&travelmode=walking'),
            ('text', '.', None),
        ])
        break

# Update document metadata.
doc.core_properties.subject = 'ברלין, אברסוולדה, שוורין, ליבק והמבורג – אוגוסט 2026'
doc.save(DOCX)
print(f'Updated {DOCX}')
